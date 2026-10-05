"""Execute production parser functions with real XLS/XLSX files.

AST isolation avoids importing the API/server dependencies for this parser unit
suite. It does not mock parsers or claim full-service acceptance; the separate
file-intelligence CI suite imports and exercises the real file_service module.
"""
import ast
import hashlib
import struct
import tempfile
import unittest
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]


def parser_namespace():
    source = ast.parse((ROOT / 'file_intelligence/file_service.py').read_text())
    names = {'_text_from_xlsx', '_text_from_xls', '_cache_key'}
    module = ast.Module(body=[n for n in source.body if isinstance(n, ast.FunctionDef) and n.name in names], type_ignores=[])
    namespace = {'Path': Path, 'Any': Any, 'hashlib': hashlib,
                 'MAX_SPREADSHEET_CELLS': 50000, 'MAX_XLS_ROWS': 10000,
                 'MAX_ARCHIVE_ENTRIES': 5000, 'MAX_ARCHIVE_EXPANDED': 512*1024*1024,
                 'MAX_ARCHIVE_TEXT_MEMBER_BYTES': 8*1024*1024, 'MAX_ARCHIVE_TEXT_TOTAL_BYTES': 32*1024*1024,
                 'MAX_ARCHIVE_LISTING_CHARS': 40000, 'ARCHIVE_LISTING_PERCENT': 25}
    exec(compile(module, 'file_intelligence/file_service.py', 'exec'), namespace)
    return namespace


def write_xls(path, rows, columns=1):
    # A genuine BIFF2 worksheet read by xlrd (no fake workbook/parser).
    def record(kind, data=b''):
        return struct.pack('<HH', kind, len(data)) + data
    data = bytearray(record(9, struct.pack('<HH', 2, 16)))
    data += record(0x42, struct.pack('<H', 1252))
    for row in range(rows):
        for column in range(columns):
            value = f'r{row}c{column}'.encode('ascii')
            data += record(4, struct.pack('<HH3sB', row, column, b'\0\0\0', len(value)) + value)
    data += record(10)
    path.write_bytes(data)


class SpreadsheetBudgets(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)
        self.parser = parser_namespace()

    def test_real_xlsx_exact_cell_budget_and_raise(self):
        from openpyxl import Workbook
        path = self.folder / 'cells.xlsx'
        book = Workbook()
        for row in range(4): book.active.append([f'r{row}c{column}' for column in range(3)])
        book.save(path); book.close()
        self.parser['MAX_SPREADSHEET_CELLS'] = 5
        text, metadata = self.parser['_text_from_xlsx'](path)
        self.assertEqual(metadata['cells_read'], 5)
        self.assertTrue(metadata['output_truncated'])
        self.assertNotIn('r1c2', text)
        self.parser['MAX_SPREADSHEET_CELLS'] = 12
        text, metadata = self.parser['_text_from_xlsx'](path)
        self.assertEqual(metadata['cells_read'], 12)
        self.assertFalse(metadata['output_truncated'], 'exact-fit workbook falsely truncated')
        self.assertIn('r3c2', text)

    def test_real_xlsx_owner_can_raise_above_old_cell_ceiling(self):
        from openpyxl import Workbook
        path = self.folder / 'large.xlsx'
        book = Workbook(write_only=True); sheet = book.create_sheet('data')
        for row in range(50001): sheet.append([f'row{row}'])
        book.save(path); book.close()
        self.parser['MAX_SPREADSHEET_CELLS'] = 50001
        text, metadata = self.parser['_text_from_xlsx'](path)
        self.assertEqual(metadata['cells_read'], 50001)
        self.assertFalse(metadata['output_truncated'])
        self.assertIn('row50000', text)

    def test_real_xlsx_shared_cell_budget_across_sheets(self):
        from openpyxl import Workbook
        path = self.folder / 'sheets.xlsx'
        book = Workbook(); book.active.append(['first', 'second'])
        book.create_sheet('later').append(['third', 'fourth']); book.save(path); book.close()
        self.parser['MAX_SPREADSHEET_CELLS'] = 3
        text, metadata = self.parser['_text_from_xlsx'](path)
        self.assertEqual(metadata['cells_read'], 3)
        self.assertTrue(metadata['output_truncated'])
        self.assertIn('third', text); self.assertNotIn('fourth', text)

    def test_real_xls_row_budget_reports_omission_and_owner_can_raise_above_old_ceiling(self):
        path = self.folder / 'rows.xls'; write_xls(path, 10001)
        text, metadata = self.parser['_text_from_xls'](path)
        self.assertTrue(metadata['output_truncated'])
        self.assertIn('xls_rows', metadata['truncation_reasons'])
        self.assertEqual(metadata['sheets'][0]['rows_read'], 10000)
        self.assertNotIn('r10000c0', text)
        self.parser['MAX_XLS_ROWS'] = 10001
        text, metadata = self.parser['_text_from_xls'](path)
        self.assertFalse(metadata['output_truncated'])
        self.assertIn('r10000c0', text)

    def test_real_xls_partial_row_respects_exact_cell_budget(self):
        path = self.folder / 'cells.xls'; write_xls(path, 4, 3)
        self.parser['MAX_SPREADSHEET_CELLS'] = 5
        text, metadata = self.parser['_text_from_xls'](path)
        self.assertEqual(metadata['cells_read'], 5)
        self.assertIn('spreadsheet_cells', metadata['truncation_reasons'])
        self.assertNotIn('r1c2', text)

    def test_budget_changes_invalidate_actual_production_cache_key(self):
        path = self.folder / 'cache.xls'; write_xls(path, 2)
        key = self.parser['_cache_key'](path, '', False, 160000)
        self.parser['MAX_XLS_ROWS'] = 10001
        changed = self.parser['_cache_key'](path, '', False, 160000)
        self.assertNotEqual(key, changed)
        self.parser['MAX_SPREADSHEET_CELLS'] = 50001
        self.assertNotEqual(changed, self.parser['_cache_key'](path, '', False, 160000))


if __name__ == '__main__':
    unittest.main()
